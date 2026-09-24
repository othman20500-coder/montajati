#!/usr/bin/env python3
"""نقل أحكام المراجعة إلى قاعدة المنصة. يعمل أكثر من مرة دون تكرار.

المصدران:
  --artifact-dir  مجلد تصدير صفحة المراجعة الحالية، بالبنية:
                  <dir>/reviewers/<uid>.json  <dir>/approvals/<uid>.json  <dir>/reviews/<uid>.json
                  (يُنتجه Claude بأداة ArtifactData، action=list مع out_dir)
  --excel         استمارة Excel أو أكثر (andalus_review_form.xlsx) معادة من مراجعين خارجيين.

  --owner         اسم صاحب المشروع، يُسجَّل صاحبًا لقرارات الاعتماد المنقولة من صفحة المراجعة.

مراجعو Excel يُسجَّلون دون اعتماد؛ صاحب المشروع يعتمدهم في المنصة، فتدخل أحكامهم في الحساب.

python tools/import_reviews.py --db andalus.sqlite --schema spec/review/schema_review.sql \
    --artifact-dir export/ --excel form_ali.xlsx form_sara.xlsx --owner "عثمان الشهري"
"""
import argparse, glob, hashlib, json, os, sqlite3, sys

VERDICT = {'ok': 'مطابق', 'minor': 'مطابق بفرق يسير', 'bad': 'غير مطابق',
           'مطابق': 'مطابق', 'مطابق بفرق يسير': 'مطابق بفرق يسير', 'غير مطابق': 'غير مطابق'}


def load_dir(d):
    out = {}
    for f in glob.glob(os.path.join(d, '*.json')):
        doc = json.load(open(f, encoding='utf-8'))
        out[os.path.splitext(os.path.basename(f))[0]] = doc.get('data', doc)
    return out


def upsert_reviewer(con, rid, source, p, declaration=None):
    con.execute("""INSERT INTO reviewers(reviewer_id,source,full_name,degree,field,org,langs,declaration,registered_at)
                   VALUES(?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(reviewer_id) DO UPDATE SET full_name=excluded.full_name, degree=excluded.degree,
                     field=excluded.field, org=excluded.org, langs=excluded.langs,
                     declaration=COALESCE(excluded.declaration, reviewers.declaration)""",
                (rid, source, p.get('name', '').strip(), p.get('degree'), p.get('field', '').strip(),
                 p.get('org'), p.get('langs'), declaration, p.get('at') or ''))


def upsert_review(con, att, rid, verdict, note, at, source, stats):
    if not con.execute('SELECT 1 FROM attestations WHERE attestation_id=?', (att,)).fetchone():
        stats['unknown_attestation'].append(att); return
    cur = con.execute('SELECT verdict, note_ar, reviewed_at FROM attestation_reviews WHERE attestation_id=? AND reviewer_id=?', (att, rid)).fetchone()
    if cur and (cur[2] or '') >= (at or ''):
        stats['unchanged'] += 1; return
    con.execute("""INSERT INTO attestation_reviews(attestation_id,reviewer_id,verdict,note_ar,reviewed_at,source)
                   VALUES(?,?,?,?,?,?)
                   ON CONFLICT(attestation_id,reviewer_id) DO UPDATE SET verdict=excluded.verdict,
                     note_ar=excluded.note_ar, reviewed_at=excluded.reviewed_at, source=excluded.source""",
                (att, rid, verdict, note or None, at, source))
    stats['written'] += 1


def import_artifact(con, d, owner, stats):
    reviewers = load_dir(os.path.join(d, 'reviewers'))
    approvals = load_dir(os.path.join(d, 'approvals'))
    reviews = load_dir(os.path.join(d, 'reviews'))
    for uid, p in reviewers.items():
        upsert_reviewer(con, 'art:' + uid, 'artifact', p); stats['reviewers'] += 1
    for uid, a in approvals.items():
        rid = 'art:' + uid
        if uid not in reviewers:
            stats['approval_without_reviewer'].append(uid); continue
        last = con.execute('SELECT approved FROM reviewer_approvals WHERE reviewer_id=? ORDER BY decided_at DESC, approval_id DESC LIMIT 1', (rid,)).fetchone()
        want = 1 if a.get('approved') else 0
        if not last or last[0] != want:
            con.execute('INSERT OR IGNORE INTO reviewer_approvals(reviewer_id,approved,decided_by,decided_at,note_ar) VALUES(?,?,?,?,?)',
                        (rid, want, owner, a.get('at') or '', 'منقول من صفحة المراجعة'))
    for uid, doc in reviews.items():
        rid = 'art:' + uid
        if uid not in reviewers:
            stats['review_without_reviewer'].append(uid); continue
        for att, x in (doc.get('items') or {}).items():
            v = VERDICT.get(x.get('v') or '')
            if v:
                upsert_review(con, att, rid, v, (x.get('note') or '').strip(), x.get('at') or '', 'artifact', stats)


def import_excel(con, path, stats):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    r = wb['المراجع']
    info = {r.cell(i, 1).value: r.cell(i, 2).value for i in range(4, 12)}
    p = {'name': str(info.get('الاسم الكامل') or '').strip(), 'degree': info.get('الدرجة العلمية'),
         'field': str(info.get('التخصص الدقيق') or '').strip(), 'org': info.get('الجهة'),
         'langs': info.get('اللغات التي يقرأ بها المصادر'), 'at': str(info.get('تاريخ المراجعة') or '')}
    if not p['name'] or not p['field'] or not str(info.get('التوقيع') or '').strip():
        stats['excel_rejected'].append(f'{os.path.basename(path)}: بيانات المراجع أو التوقيع ناقصة'); return
    rid = 'xl:' + hashlib.sha1((p['name'] + '|' + p['field']).encode()).hexdigest()[:16]
    upsert_reviewer(con, rid, 'excel', p, declaration=str(info.get('الإقرار') or '')); stats['reviewers'] += 1
    s = wb['الشواهد']
    hdr = [c.value for c in s[1]]
    ci = {h: i for i, h in enumerate(hdr)}
    for row in s.iter_rows(min_row=2, values_only=True):
        att, v = row[ci['المعرف']], VERDICT.get(str(row[ci['الحكم']] or '').strip())
        if att and v:
            upsert_review(con, att, rid, v, str(row[ci['الملاحظة أو النص الصحيح']] or '').strip(), p['at'], 'excel', stats)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--db', required=True); ap.add_argument('--schema', required=True)
    ap.add_argument('--artifact-dir'); ap.add_argument('--excel', nargs='*', default=[])
    ap.add_argument('--owner', default='صاحب المشروع')
    a = ap.parse_args()
    con = sqlite3.connect(a.db); con.execute('PRAGMA foreign_keys=ON')
    con.executescript(open(a.schema, encoding='utf-8').read())
    stats = {'reviewers': 0, 'written': 0, 'unchanged': 0, 'unknown_attestation': [], 'approval_without_reviewer': [],
             'review_without_reviewer': [], 'excel_rejected': []}
    with con:
        if a.artifact_dir: import_artifact(con, a.artifact_dir, a.owner, stats)
        for x in a.excel: import_excel(con, x, stats)
    stats['status_counts'] = dict(con.execute('SELECT review_status, COUNT(*) FROM v_attestation_review_status GROUP BY 1').fetchall())
    print(json.dumps(stats, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
