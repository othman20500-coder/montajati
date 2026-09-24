#!/usr/bin/env python3
"""اختبار خاصية المراجعة ببيانات تجريبية (أسماء المراجعين وهمية).
يبني نسخة مؤقتة من seed، ويشغّل سكربت النقل مرتين، ثم فحوص المراجعة، ثم حالات يجب أن تفشل.

python tests/test_review_feature.py
"""
import json, os, shutil, sqlite3, subprocess, sys, tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SEED = os.path.join(ROOT, 'seed', 'andalus_atlas_v1_2.sqlite')
SCHEMA = os.path.join(ROOT, 'spec', 'review', 'schema_review.sql')
sys.path.insert(0, os.path.join(ROOT, 'tests'))
import andalus_review_qa as qa  # noqa: E402

A1, A2, A3 = 'ATT-004', 'ATT-005', 'ATT-263'


def write(p, obj):
    os.makedirs(os.path.dirname(p), exist_ok=True)
    json.dump(obj, open(p, 'w', encoding='utf-8'), ensure_ascii=False)


def make_export(d):
    write(f'{d}/reviewers/u1.json', {'name': 'مراجع تجريبي أ', 'degree': 'دكتوراه', 'field': 'تاريخ الأندلس', 'org': 'جامعة تجريبية', 'at': '2026-09-24T10:00:00Z'})
    write(f'{d}/reviewers/u2.json', {'name': 'مراجع تجريبي ب', 'degree': 'دكتوراه', 'field': 'تحقيق المخطوطات', 'at': '2026-09-24T10:05:00Z'})
    write(f'{d}/reviewers/u3.json', {'name': 'مراجع غير معتمد', 'degree': 'بكالوريوس', 'field': 'تاريخ', 'at': '2026-09-24T10:06:00Z'})
    write(f'{d}/approvals/u1.json', {'approved': True, 'at': '2026-09-24T11:00:00Z'})
    write(f'{d}/approvals/u2.json', {'approved': True, 'at': '2026-09-24T11:00:00Z'})
    write(f'{d}/reviews/u1.json', {'items': {A1: {'v': 'ok', 'note': '', 'at': '2026-09-24T12:00:00Z'},
                                             A2: {'v': 'ok', 'note': '', 'at': '2026-09-24T12:01:00Z'},
                                             A3: {'v': 'minor', 'note': 'فرق في الهمزة', 'at': '2026-09-24T12:02:00Z'}}})
    write(f'{d}/reviews/u2.json', {'items': {A1: {'v': 'ok', 'note': '', 'at': '2026-09-24T12:10:00Z'}}})


def run_import(db, export):
    r = subprocess.run([sys.executable, os.path.join(ROOT, 'tools', 'import_reviews.py'), '--db', db, '--schema', SCHEMA,
                        '--artifact-dir', export, '--owner', 'صاحب المشروع (تجريبي)'], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    return json.loads(r.stdout)


def main():
    results = []
    def check(name, cond):
        results.append((name, bool(cond)))
    with tempfile.TemporaryDirectory() as tmp:
        db = os.path.join(tmp, 't.sqlite'); shutil.copy(SEED, db)
        exp = os.path.join(tmp, 'export'); make_export(exp)
        s1 = run_import(db, exp); s2 = run_import(db, exp)
        con = sqlite3.connect(db)
        st = dict(con.execute('SELECT attestation_id, review_status FROM v_attestation_review_status WHERE attestation_id IN (?,?,?)', (A1, A2, A3)).fetchall())
        check('نقل الأحكام: 4 أحكام', s1['written'] == 4)
        check('التشغيل الثاني لا يكرر شيئًا', s2['written'] == 0 and con.execute('SELECT COUNT(*) FROM attestation_reviews').fetchone()[0] == 4)
        check('مراجعان بالمطابقة ← معتمد بالعين', st[A1] == 'معتمد بالعين')
        check('مراجع واحد ← يحتاج مراجعًا ثانيًا', st[A2] == 'يحتاج مراجعًا ثانيًا')
        check('فرق يسير ← خلاف أو فرق', st[A3] == 'خلاف أو فرق')
        check('فحوص المراجعة تنجح على بيانات سليمة', qa.run(db)['status'] == 'PASS')

        # سحب اعتماد مراجع يُسقط أحكامه من الحساب
        con.execute("INSERT INTO reviewer_approvals(reviewer_id,approved,decided_by,decided_at) VALUES('art:u2',0,'صاحب المشروع (تجريبي)','2026-09-24T13:00:00Z')"); con.commit()
        check('سحب الاعتماد يُسقط الحكم', con.execute('SELECT review_status FROM v_attestation_review_status WHERE attestation_id=?', (A1,)).fetchone()[0] == 'يحتاج مراجعًا ثانيًا')

        def fails(name, sql, expect):
            c2 = os.path.join(tmp, name + '.sqlite'); shutil.copy(db, c2)
            k = sqlite3.connect(c2); k.executescript(sql); k.commit(); k.close()
            res = {c['check']: c['pass'] for c in qa.run(c2)['checks']}
            check('يُكتشف: ' + name, res[expect] is False)
        fails('حكم من مراجع غير معتمد',
              f"INSERT INTO attestation_reviews(attestation_id,reviewer_id,verdict,reviewed_at,source) VALUES('{A2}','art:u3','مطابق','2026-09-24T12:30:00Z','platform')",
              'review_by_approved_reviewer')
        fails('عدم مطابقة بلا ملاحظة',
              f"UPDATE attestation_reviews SET verdict='غير مطابق', note_ar=NULL WHERE attestation_id='{A2}'",
              'non_match_has_note')
        # استمارة Excel: تُنقل دون اعتماد، ولا تدخل الحساب حتى يعتمدها صاحب المشروع
        from openpyxl import load_workbook
        xl = os.path.join(tmp, 'form.xlsx'); wb = load_workbook(os.path.join(ROOT, 'tools', 'andalus_review_form.xlsx'))
        r = wb['المراجع']
        for i, v in {4: 'مراجع خارجي تجريبي', 5: 'دكتوراه', 6: 'تاريخ الأندلس', 9: '2026-09-25', 11: 'موقّع'}.items(): r.cell(i, 2, v)
        sh = wb['الشواهد']
        for row in range(2, sh.max_row + 1):
            if sh.cell(row, 1).value == A2: sh.cell(row, 9, 'مطابق')
        wb.save(xl)
        rr = subprocess.run([sys.executable, os.path.join(ROOT, 'tools', 'import_reviews.py'), '--db', db, '--schema', SCHEMA, '--excel', xl], capture_output=True, text=True)
        sx = json.loads(rr.stdout)
        check('استمارة Excel تُنقل', sx['written'] == 1 and not sx['excel_rejected'])
        check('حكم Excel لا يُحتسب قبل الاعتماد', con.execute('SELECT review_status FROM v_attestation_review_status WHERE attestation_id=?', (A2,)).fetchone()[0] == 'يحتاج مراجعًا ثانيًا')
        xid = con.execute("SELECT reviewer_id FROM reviewers WHERE source='excel'").fetchone()[0]
        con.execute("INSERT INTO reviewer_approvals(reviewer_id,approved,decided_by,decided_at) VALUES(?,1,'صاحب المشروع (تجريبي)','2026-09-25T09:00:00Z')", (xid,)); con.commit()
        check('بعد اعتماده يصير الشاهد معتمدًا بالعين', con.execute('SELECT review_status FROM v_attestation_review_status WHERE attestation_id=?', (A2,)).fetchone()[0] == 'معتمد بالعين')
        check('فحوص المراجعة تنجح بعد نقل Excel', qa.run(db)['status'] == 'PASS')
        fails('وصف «مطابق بالعين» بلا مراجعَين',
              f"UPDATE attestations SET verbatim_check='مطابق بالعين (مراجعة تجريبية)' WHERE attestation_id='{A3}'",
              'eye_verified_label_backed')
    ok = sum(r for _, r in results)
    print(json.dumps({'status': 'PASS' if ok == len(results) else 'FAIL', 'passed': ok, 'total': len(results),
                      'tests': [{'test': n, 'pass': r} for n, r in results]}, ensure_ascii=False, indent=1))
    sys.exit(0 if ok == len(results) else 1)


if __name__ == '__main__':
    main()
